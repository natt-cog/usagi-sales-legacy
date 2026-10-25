#!/usr/bin/env python3
"""spec/設計書.xlsx（静的）を複製し、判定・判定の根拠・エビデンス列と、修正箇所一覧・要確認の挙動・業務ルール検証結果のシートを足して
spec/dynamic/設計書.xlsx に保存する。判定結果は evidence/verdicts.json にも書き出し、build_html.py が使う。"""
import os, json, re, copy
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from verdicts import *
HERE = os.path.dirname(os.path.abspath(__file__)); DYN = os.path.abspath(os.path.join(HERE, '..', '..')); SPEC = os.path.dirname(DYN)
J = os.path.join
rules = {x['row']: x for f in ('batch_results.json', 'web_results.json') for x in json.load(open(J(DYN, 'evidence', 'rules', f)))}
exs = {x['row']: x for x in json.load(open(J(DYN, 'evidence', 'examples', 'examples.json')))}
crud = json.load(open(J(DYN, 'evidence', 'crud', 'crud_actual.json')))
wb = openpyxl.load_workbook(J(SPEC, '設計書.xlsx'))
FILL = {EXEC: 'FFC6EFCE', SRC: 'FFDDEBF7', EST: 'FFFFEB9C', NOT: 'FFF4CCCC'}
HDR = PatternFill('solid', fgColor='FFDDEBF7'); FIXFILL = PatternFill('solid', fgColor='FFFFD966')
WRAP = Alignment(wrap_text=True, vertical='top')
out = {'program_spec': {}, 'program_list': {}, 'crud': {}, 'diff': {}, 'questions': {}}
fixes = []


def add_cols(ws, names, widths):
    c0 = ws.max_column + 1
    for i, (n, w) in enumerate(zip(names, widths)):
        c = ws.cell(1, c0 + i, n); c.font = Font(bold=True); c.fill = HDR; c.alignment = WRAP
        ws.column_dimensions[openpyxl.utils.get_column_letter(c0 + i)].width = w
    return c0


def put(ws, r, c0, verdict, why, ev, extra=()):
    a = ws.cell(r, c0, verdict); a.fill = PatternFill('solid', fgColor=FILL[verdict]); a.alignment = WRAP
    ws.cell(r, c0 + 1, why).alignment = WRAP
    if ev:
        e = ws.cell(r, c0 + 2, ev); e.hyperlink = ev; e.font = Font(color='FF0563C1', underline='single'); e.alignment = WRAP
    for i, v in enumerate(extra):
        ws.cell(r, c0 + 3 + i, v).alignment = WRAP


# ---- プログラム仕様書 ----
ws = wb['プログラム仕様書']
c0 = add_cols(ws, ['判定', '判定の根拠（入力→期待結果→実際の結果）', 'エビデンス'], [12, 70, 38])
fixmatch = {f['id']: f for f in FIXES if 'match' in f}
for r in range(2, ws.max_row + 1):
    pid, kubun = ws.cell(r, 1).value, ws.cell(r, 4).value
    tgt, item, body = (ws.cell(r, 5).value or ''), (ws.cell(r, 6).value or ''), (ws.cell(r, 7).value or '')
    text = f'{tgt} {item} {body}'
    if kubun == '業務ルール':
        x = rules[r]
        v = EXEC if x['result'] == '一致' else EXEC
        why = f"[{x['id']}] 入力: {x['input']} → 期待: {x['expected']} → 実際: {x['actual']}（{x['result']}）" + (f" ※{x['note']}" if x['note'] else '')
        ev = x['evidence']
    elif kubun in ('入力例', '出力例'):
        x = exs.get(r)
        if x:
            v = x['verdict']; why = ('実際: ' if v == EXEC else '') + x['actual'] + ('' if x['ok'] in (True, None) else '（設計書と不一致 → 修正）'); ev = x['evidence']
            if x['fix'] and v == EXEC:
                f = next((f for f in FIXES if f.get('row') == r), None)
                kind = f['kind'] if f else '推定→確定'
                fixes.append(dict(id=f['id'] if f else '', kind=kind, doc='01 プログラム仕様書', where=f'{pid} {ws.cell(r, 3).value} {kubun}「{tgt}」（Excel {r}行）',
                                  before=body, after=x['fix'], evidence=ev, reason=f['reason'] if f else '静的版は推定・値不明。実行結果で確定'))
                ws.cell(r, 7, x['fix']).fill = FIXFILL
        else:
            v, why, ev = SRC, '値はソース・データから読んだもの（実行では使っていない）', ''
    elif kubun == '概要':
        v, why, ev = EXEC, 'プログラムを実行し、概要の各文を業務ルール検証で確認（同プログラムの業務ルール行を参照）', PEV[pid]
    else:  # 入力レコード／出力レコード
        if 'NYKREC' in text or (pid == 'P17' and '入金' in tgt and '推定' in text):
            v, why, ev = EXEC, '復元案の NYKREC で実データ650件を読み、消込結果と請求書サンプルが一致（定義そのものは復元案）', E + 'batch/run_202412/USAGI.PROD.NYUKIN'
        elif '推定' in text:
            v, why, ev = EST, '静的版の推定のまま（実行で値が観測できない項目）', ''
        elif 'FILLER' in text or '参照しない' in text or '使わない' in text:
            v, why, ev = SRC, 'プログラムが参照しない項目のため実行結果に影響しない', ''
        elif 'セッション' in tgt:
            v, why, ev = SRC, 'セッションの中身は画面から直接見えない（ログイン判定の動作のみ実行で確認）', PEV.get(pid, '')
        elif pid in ('P13', 'P14', 'P15', 'P16', 'P17', 'P18', 'P19', 'P20', 'P21'):
            v = EXEC; why = '固定長の桁位置で実データを読み書きし、帳票サンプルと一致。項目値は境界値試験の入出力で確認'
            ev = E + 'batch/run_202412/' + BATCH_FILE_EV[pid] if BATCH_FILE_EV.get(pid) else PEV[pid]
        elif pid in ('P11', 'P12'):
            v, why, ev = EXEC, '出力ファイル・テーブルをバイト位置で確認', PEV[pid]
        else:
            v, why, ev = EXEC, '画面操作・SQL ログで確認', (E + 'screens/' + SCREENS[pid][0]) if pid in SCREENS else PEV[pid]
    for fid, f in fixmatch.items():
        if pid == f['match'][0] and f['before'] in body:
            ws.cell(r, 7, body.replace(f['before'], f['after'])).fill = FIXFILL
            fixes.append(dict(id=fid, kind=f['kind'], doc=f['doc'], where=f"{f['where']}（Excel {r}行）", before=body, after=body.replace(f['before'], f['after']), evidence=f['evidence'], reason=f['reason']))
            v, ev = EXEC, f['evidence']; why = f['reason'] + '（修正済み）'
    put(ws, r, c0, v, why, ev)
    out['program_spec'][r] = dict(pid=pid, kubun=kubun, verdict=v, why=why, ev=ev, body=ws.cell(r, 7).value, target=tgt, item=item,
                                  name=ws.cell(r, 3).value, prog=ws.cell(r, 2).value, src=ws.cell(r, 8).value)
# ---- プログラム一覧 ----
ws = wb['プログラム一覧']; c0 = add_cols(ws, ['判定', '判定の根拠', 'エビデンス'], [12, 50, 38])
for r in range(2, ws.max_row + 1):
    pid = ws.cell(r, 1).value
    v = EXEC; why = '実行できた（画面は docker compose、COBOL は GnuCOBOL 3.1.2 でコンパイルして実行）'
    if pid == 'P17': why = 'NYKREC.cpy が欠けておりそのままではコンパイル不可。復元案を補ってコンパイル・実行できた'
    if pid == 'P21': why = 'JCL からは呼ばれない。DD URINEW を補って単独実行'
    put(ws, r, c0, v, why, PEV[pid]); out['program_list'][pid] = dict(verdict=v, why=why, ev=PEV[pid])
# ---- CRUD図 ----
ws = wb['CRUD図']; c0 = add_cols(ws, ['判定', '実測（SQL ログ・strace）', 'エビデンス'], [12, 50, 38])
heads = [c.value for c in ws[1]][2:21]
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, 1).value; static = {heads[i]: ws.cell(r, 3 + i).value for i in range(len(heads)) if ws.cell(r, 3 + i).value}
    key = name.split('/')[-1].strip().replace('（未使用）', '')
    if key in crud:
        act = {k.replace('.txt', ''): v for k, v in crud[key].items()}
        st2 = {k: v.replace('D', '') if k in ('JUCHSRT', 'NYUKSRT') else v for k, v in static.items()}
        ok = st2 == act
        why = f"実測 {act}" + ('' if ok else f' ／ 静的 {static}') + ('（D は JCL の DISP=(OLD,DELETE)。pipeline.py で削除）' if 'D' in ''.join(static.get(k, '') for k in ('JUCHSRT', 'NYUKSRT')) else '')
        ev = E + 'crud/' + (key + '.sql.txt' if not key.startswith('US') else 'strace/' + key + '.strace')
        v = EXEC
    elif key.startswith('USB900'):
        v, why, ev = EXEC, '境界値試験で URIAGE を読み URINEW を作成', E + 'rules/USB900'
    else:
        v, why, ev = EXEC, 'pipeline.py で同じ入出力を実行（FTP は docker compose exec、SORT は Python で代替）', E + 'batch/run_202412.log'
    put(ws, r, c0, v, why, ev); out['crud'][name] = dict(verdict=v, why=why, ev=ev)
# ---- 旧仕様書との差分 ----
ws = wb['旧仕様書との差分']; c0 = add_cols(ws, ['判定', '判定の根拠', 'エビデンス'], [12, 40, 38])
for r in range(2, ws.max_row + 1):
    v, ev = D[ws.cell(r, 1).value]; why = '現行の内容を実行で確認'
    put(ws, r, c0, v, why, ev); out['diff'][ws.cell(r, 1).value] = dict(verdict=v, why=why, ev=ev)
# ---- 確認事項一覧 ----
ws = wb['確認事項一覧']; c0 = add_cols(ws, ['判定', '実行結果', 'エビデンス'], [12, 50, 38])
for r in range(2, ws.max_row + 1):
    q = ws.cell(r, 1).value; v, why, ev = Q[q]
    if q == 'Q03':
        old = ws.cell(r, 3).value; new = re.sub(r'（推定：[^）]*）', '（実行で確認）', old)
        ws.cell(r, 3, new).fill = FIXFILL
    put(ws, r, c0, v, why, ev); out['questions'][q] = dict(verdict=v, why=why, ev=ev)
# ---- 追加シート ----
fixes += [dict(f, reason=f.get('reason', '静的版は推定。実行で確定')) for f in FIXES_MD]
order = {'誤り': 0, '不正確': 1, '推定→確定': 2}
fixes.sort(key=lambda f: (order[f['kind']], f['id'] or 'F5'))
n = 0
for f in fixes:
    if not f['id']:
        n += 1; f['id'] = f'F{10 + n:02d}'
def sheet(name, head, rows_, widths):
    w = wb.create_sheet(name)
    w.append(head)
    for c in w[1]: c.font = Font(bold=True); c.fill = HDR; c.alignment = WRAP
    for row in rows_:
        w.append([row[i] for i in range(len(head))])
    for i, wd in enumerate(widths): w.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = wd
    for rr in w.iter_rows(min_row=2):
        for c in rr: c.alignment = WRAP
        ev = rr[-1]
        if ev.value and str(ev.value).startswith('evidence/'):
            ev.hyperlink = ev.value; ev.font = Font(color='FF0563C1', underline='single')
    w.freeze_panes = 'A2'; w.auto_filter.ref = w.dimensions
sheet('修正箇所一覧', ['No', '種類', '設計書', '箇所', '修正前', '修正後', '理由', 'エビデンス'],
      [[f['id'], f['kind'], f['doc'], f['where'], f['before'], f['after'], f['reason'], f['evidence']] for f in fixes], [6, 10, 18, 36, 50, 50, 40, 38])
sheet('要確認の挙動', ['No', '挙動', '内容（実行結果）', '関連する確認事項', 'エビデンス'], ISSUES, [6, 34, 70, 12, 38])
allr = sorted(rules.values(), key=lambda x: x['row'])
sheet('業務ルール検証結果', ['仕様書の行', 'プログラム', 'ケース', '入力', '期待結果（設計書から）', '実際の結果', '結果', '備考', 'エビデンス'],
      [[x['row'], x['prog'], x['id'], x['input'], x['expected'], x['actual'], x['result'], x['note'], x['evidence']] for x in allr], [8, 16, 14, 40, 40, 50, 6, 30, 30])
# 集計
cnt = {}
for sec in ('program_spec', 'program_list', 'crud', 'diff', 'questions'):
    for v in out[sec].values(): cnt[v['verdict']] = cnt.get(v['verdict'], 0) + 1
out['fixes'] = fixes; out['count'] = cnt; out['issues'] = ISSUES
w = wb.create_sheet('判定の集計', 0)
w.append(['判定', '件数', '割合']); [setattr(c, 'font', Font(bold=True)) for c in w[1]]
tot = sum(cnt.values())
for k in (EXEC, SRC, EST, NOT):
    w.append([k, cnt.get(k, 0), f'{cnt.get(k, 0) / tot:.1%}'])
w.append(['合計', tot, '100%']); w.append([])
w.append(['対象: プログラム一覧・プログラム仕様書・CRUD図・旧仕様書との差分・確認事項一覧の各行（Markdown のみの記述は index.html で判定）'])
w.append([f"修正箇所 {len(fixes)} 件（誤り {sum(f['kind']=='誤り' for f in fixes)}・不正確 {sum(f['kind']=='不正確' for f in fixes)}・推定→確定 {sum(f['kind']=='推定→確定' for f in fixes)}）／要確認の挙動 {len(ISSUES)} 件"])
w.column_dimensions['A'].width = 14
wb.save(J(DYN, '設計書.xlsx'))
json.dump(out, open(J(DYN, 'evidence', 'verdicts.json'), 'w'), ensure_ascii=False, indent=1)
print(cnt, tot, 'fixes', len(fixes), [(f['id'], f['kind'], f['where'][:40]) for f in fixes if f['kind'] != '推定→確定'])
